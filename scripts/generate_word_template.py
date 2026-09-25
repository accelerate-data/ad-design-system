# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx==1.2.0", "fonttools[woff]==4.66.0"]
# ///
"""Build the corporate Word template. Run: uv run scripts/generate_word_template.py."""

from io import BytesIO
from pathlib import Path
import re
import uuid
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from fontTools.ttLib import TTFont
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
DESIGN = ROOT / "plugin/skills/applying-design-system"
OUT = ROOT / "templates"
STEM = "Accelerate-Data-Document-Template"
TOKENS = (DESIGN / "tokens/colors.css").read_text()


def color(name):
    return re.search(rf"--{name}:\s*#([0-9a-fA-F]{{6}})", TOKENS)[1].upper()


def element(tag, **attrs):
    item = OxmlElement(tag)
    for key, value in attrs.items():
        item.set(qn(key), str(value))
    return item


def font(style, family, size, ink, bold=False):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color(ink))
    style.font.bold = bold
    style.font.italic = False
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if "theme" in key.lower():
            del fonts.attrib[key]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), family)


def border(paragraph, side="bottom", ink="pacific", size=8, space=8):
    props = paragraph._p.get_or_add_pPr()
    borders = element("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        if edge != side:
            borders.append(element(f"w:{edge}", **{"w:val": "nil"}))
    borders.append(element(f"w:{side}", **{
        "w:val": "single", "w:sz": size, "w:space": space, "w:color": color(ink)
    }))
    props.append(borders)


def field(paragraph, code, fallback):
    run = paragraph.add_run()
    run._r.append(element("w:fldChar", **{"w:fldCharType": "begin"}))
    instruction = element("w:instrText", **{"xml:space": "preserve"})
    instruction.text = f" {code} "
    run._r.append(instruction)
    run._r.append(element("w:fldChar", **{"w:fldCharType": "separate"}))
    paragraph.add_run(fallback)
    paragraph.add_run()._r.append(element("w:fldChar", **{"w:fldCharType": "end"}))


def table_style(document):
    style = document.styles.add_style("AD table", WD_STYLE_TYPE.TABLE)
    font(style, "Geist", 10, "smoke")
    style.paragraph_format.space_after = Pt(3)
    props = element("w:tblPr")
    margins = element("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        margins.append(element(f"w:{edge}", **{"w:w": 120, "w:type": "dxa"}))
    props.append(margins)
    borders = element("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element(f"w:{edge}", **{
            "w:val": "single", "w:sz": 4, "w:color": color("gray-200")
        }))
    props.append(borders)
    style.element.append(props)
    first = element("w:tblStylePr", **{"w:type": "firstRow"})
    cell = element("w:tcPr")
    cell.append(element("w:shd", **{"w:fill": color("navy")}))
    run = element("w:rPr")
    run.append(element("w:b"))
    run.append(element("w:color", **{"w:val": color("white")}))
    first.append(run)
    first.append(cell)
    style.element.append(first)


def build():
    doc = Document()
    props = doc.core_properties
    props.title = "[Document title]"
    props.subject = "Accelerate Data document template"
    props.author = "Accelerate Data"
    props.keywords = "Accelerate Data, corporate, template"
    props.comments = ""
    styles = doc.styles
    font(styles["Normal"], "Geist", 10.5, "smoke")
    styles["Normal"].paragraph_format.line_spacing = 1.4
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Normal"].paragraph_format.widow_control = True
    for name, size in (("Title", 36), ("Heading 1", 22), ("Heading 2", 15), ("Heading 3", 12)):
        font(styles[name], "Geist SemiBold", size, "navy")
        styles[name].paragraph_format.space_before = Pt(18 if name != "Title" else 0)
        styles[name].paragraph_format.space_after = Pt(8)
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.keep_together = True
        styles[name].next_paragraph_style = styles["Normal"]
    for name in ("Subtitle", "Caption", "Header", "Footer"):
        font(styles[name], "Geist", 9 if name != "Subtitle" else 15, "gray-600")
        styles[name].paragraph_format.space_after = Pt(6)
    for name in ("Title", "Subtitle"):
        for node in list(styles[name].element.get_or_add_pPr()):
            if node.tag in (qn("w:pBdr"), qn("w:numPr")):
                node.getparent().remove(node)
        for node in list(styles[name].element.get_or_add_rPr()):
            if node.tag == qn("w:spacing"):
                node.getparent().remove(node)
    for name in ("Header", "Footer"):
        styles[name].paragraph_format.tab_stops.clear_all()
    for name in ("List Bullet", "List Number"):
        font(styles[name], "Geist", 10.5, "smoke")
        styles[name].paragraph_format.space_after = Pt(5)
    for name, family, size, ink in (
        ("AD label", "Geist Mono", 9, "ocean"),
        ("AD callout", "Geist", 12, "navy"),
        ("AD metadata", "Geist Mono", 9, "gray-600"),
        ("AD code", "Geist Mono", 9, "smoke"),
        ("AD table header", "Geist", 10, "white"),
    ):
        s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = styles["Normal"]
        font(s, family, size, ink)
        s.next_paragraph_style = styles["Normal"]
    styles["AD label"].paragraph_format.keep_with_next = True
    styles["AD table header"].font.bold = True
    styles["AD table header"].paragraph_format.space_after = Pt(3)
    cp = styles["AD callout"].element.get_or_add_pPr()
    cp.append(element("w:shd", **{"w:fill": color("powder")}))
    styles["AD callout"].paragraph_format.space_before = Pt(12)
    styles["AD callout"].paragraph_format.space_after = Pt(16)
    styles["AD callout"].paragraph_format.keep_together = True
    styles["AD code"].element.get_or_add_pPr().append(
        element("w:shd", **{"w:fill": color("pearl")}))
    styles["AD code"].paragraph_format.keep_together = True
    font(styles["Strong"], "Geist", 10.5, "smoke", bold=True)
    table_style(doc)

    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(27), Mm(23)
    section.left_margin = section.right_margin = Mm(23)
    section.header_distance = section.footer_distance = Mm(12)
    section.different_first_page_header_footer = True
    width = section.page_width - section.left_margin - section.right_margin
    header = section.header.paragraphs[0]
    header.add_run().add_picture(str(ROOT / "logo/product/ui/logo-dark-h64@2x.png"), width=Mm(48))
    header.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    header.add_run("\t[Short document title]")
    border(header, ink="gray-200", size=4)
    footer = section.footer.paragraphs[0]
    footer.add_run("Accelerate Data  ·  [Classification]")
    footer.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    footer.add_run("\t")
    field(footer, "PAGE", "2")
    footer.add_run(" / ")
    field(footer, "NUMPAGES", "3")
    border(footer, side="top", ink="gray-200", size=4)
    section.first_page_footer.paragraphs[0].text = "acceleratedata.ai  ·  [Classification]"
    doc.settings.element.append(element("w:updateFields", **{"w:val": "true"}))

    # Cover: ordinary paragraphs remain editable and flow naturally for long titles.
    p = doc.add_paragraph()
    p.add_run().add_picture(str(ROOT / "logo/product/ui/logo-dark-h64@2x.png"), width=Mm(78))
    p.paragraph_format.space_after = Pt(100)
    doc.add_paragraph("[DOCUMENT TYPE]", "AD label")
    doc.add_paragraph("[Document title]", "Title")
    p = doc.add_paragraph("[A short subtitle that states the purpose and intended outcome.]", "Subtitle")
    p.paragraph_format.space_after = Pt(30)
    border(p, size=12, space=16)
    for label, value in (("PREPARED FOR", "[Client or audience]"), ("PREPARED BY", "[Author or team]"), ("DATE / VERSION", "[DD Month YYYY]  /  [v0.1]")):
        doc.add_paragraph(label, "AD label")
        p = doc.add_paragraph(value)
        p.paragraph_format.space_after = Pt(15)
    doc.add_page_break()

    doc.add_paragraph("OVERVIEW", "AD label")
    doc.add_heading("Executive summary", 1)
    doc.add_paragraph("[State the main recommendation or finding. Explain why it matters to this audience and what decision or action this document supports.]")
    p = doc.add_paragraph("[Key message: the one point the reader should remember.]", "AD callout")
    border(p, side="left", size=16, space=10)
    doc.add_heading("Context and objectives", 2)
    doc.add_paragraph("[Describe the current situation, the problem to solve, and the result you want to achieve.]")
    for item in ("[Objective one and its measure of success.]", "[Objective two and its measure of success.]", "[Objective three and its measure of success.]"):
        doc.add_paragraph(item, "List Bullet")
    doc.add_heading("Scope", 2)
    doc.add_paragraph("[Define what this document covers. State key boundaries, assumptions, and dependencies.]")
    doc.add_heading("Decision required", 2)
    doc.add_paragraph("[Specify the decision, the person accountable for it, and the date it is needed.]")
    doc.add_page_break()

    doc.add_paragraph("DETAIL", "AD label")
    doc.add_heading("[Section heading]", 1)
    doc.add_paragraph("[Present the evidence, approach, or proposal. Use Heading 1, Heading 2, and Heading 3 so the document remains easy to navigate and extend.]")
    doc.add_heading("[Subsection heading]", 2)
    doc.add_paragraph("[Explain the main point. Add only the supporting detail the reader needs.]")
    doc.add_heading("[Supporting detail]", 3)
    doc.add_paragraph("[Record a constraint, assumption, or technical detail.]")
    doc.add_paragraph("[reference_id]  |  [version]  |  [value]", "AD code")
    doc.add_heading("Actions and owners", 2)
    table = doc.add_table(rows=1, cols=3, style="AD table")
    table.autofit = False
    for column, mm in zip(table.columns, (88, 40, 36)):
        column.width = Mm(mm)
    for cell, label in zip(table.rows[0].cells, ("Action", "Owner", "Due date")):
        cell.text = label
        cell.paragraphs[0].style = styles["AD table header"]
    table.rows[0]._tr.get_or_add_trPr().append(element("w:tblHeader"))
    for n in range(1, 4):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        for cell, value in zip(row.cells, (f"[Action {n}]", "[Name]", "[DD Mon YYYY]")):
            cell.text = value
        row.cells[2].paragraphs[0].style = styles["AD metadata"]
    doc.add_paragraph("Table 1. [Describe the actions and their intended result.]", "Caption")
    doc.add_heading("Next steps", 2)
    doc.add_paragraph("[Confirm the decision and assign the work.]", "List Number")
    doc.add_paragraph("[Agree the next review date and success criteria.]", "List Number")
    doc.add_heading("References", 2)
    doc.add_paragraph("[Source title, author, publication date, and URL or document reference.]", "Caption")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def package(source):
    with ZipFile(BytesIO(source)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    theme = etree.fromstring(parts["word/theme/theme1.xml"])
    theme.set("name", "Accelerate Data")
    for node in theme.findall(".//a:fontScheme", ns):
        node.set("name", "Accelerate Data")
    for path, family in ((".//a:majorFont/a:latin", "Geist SemiBold"), (".//a:minorFont/a:latin", "Geist")):
        theme.find(path, ns).set("typeface", family)
    scheme = theme.find(".//a:clrScheme", ns)
    scheme.set("name", "Accelerate Data")
    for tag, ink in zip(("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"),
                        ("smoke", "white", "navy", "pearl", "pacific", "navy", "seafoam", "powder", "arctic", "ocean", "ocean", "navy")):
        node = scheme.find(f"a:{tag}", ns)
        for child in list(node):
            node.remove(child)
        etree.SubElement(node, f"{{{ns['a']}}}srgbClr", val=color(ink))
    parts["word/theme/theme1.xml"] = etree.tostring(theme)
    # Embed full font files so new text is supported, not just the sample glyphs.
    fonts = etree.fromstring(parts["word/fontTable.xml"])
    relns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rels = etree.Element(f"{{{relns}}}Relationships", nsmap={None: relns})
    for index, (file, family, variant) in enumerate((
        ("Geist-Regular.woff2", "Geist", "embedRegular"),
        ("Geist-Bold.woff2", "Geist", "embedBold"),
        ("Geist-SemiBold.woff2", "Geist SemiBold", "embedRegular"),
        ("GeistMono-Regular.woff2", "Geist Mono", "embedRegular"),
    ), 1):
        f = TTFont(DESIGN / "assets/fonts" / file, recalcTimestamp=False)
        if f["OS/2"].fsType != 0:
            raise ValueError(f"Review font embedding permissions: {file}")
        f.flavor = None
        stream = BytesIO()
        f.save(stream)
        data = bytearray(stream.getvalue())
        key = uuid.uuid5(uuid.NAMESPACE_URL, f"acceleratedata.ai/fonts/{file}")
        mask = key.bytes[::-1]
        for j in range(32):
            data[j] ^= mask[j % 16]
        target = f"fonts/font{index}.odttf"
        parts[f"word/{target}"] = bytes(data)
        node = next((n for n in fonts if n.get(qn("w:name")) == family), None)
        if node is None:
            node = element("w:font", **{"w:name": family})
            fonts.append(node)
        node.append(element(f"w:{variant}", **{
            "r:id": f"rId{index}", "w:fontKey": "{" + str(key).upper() + "}", "w:subsetted": "false"
        }))
        etree.SubElement(rels, f"{{{relns}}}Relationship", Id=f"rId{index}",
                         Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font", Target=target)
    parts["word/fontTable.xml"] = etree.tostring(fonts)
    parts["word/_rels/fontTable.xml.rels"] = etree.tostring(rels)
    settings = etree.fromstring(parts["word/settings.xml"])
    settings.append(element("w:embedTrueTypeFonts"))
    parts["word/settings.xml"] = etree.tostring(settings)
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
                     Extension="odttf", ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont")
    OUT.mkdir(exist_ok=True)
    for extension, kind in (("docx", "document"), ("dotx", "template")):
        for node in types:
            if node.get("PartName") == "/word/document.xml":
                node.set("ContentType", f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}.main+xml")
        parts["[Content_Types].xml"] = etree.tostring(types)
        path = OUT / f"{STEM}.{extension}"
        with ZipFile(path, "w", ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    package(build())
